from datetime import timedelta

import pytest
from django.utils import timezone
from rest_framework import status

from conftest import _make_jwt_for


@pytest.mark.django_db
class TestPlanViewSet:
    url = '/api/billing/plans/'

    def test_admin_can_create_plan(self, admin_client):
        resp = admin_client.post(self.url, {
            'name': 'Mensual', 'price': '1000.00', 'duration_value': 1, 'duration_unit': 'month',
        }, format='json')
        assert resp.status_code == status.HTTP_201_CREATED
        assert resp.data['name'] == 'Mensual'

    def test_list_plans(self, admin_client, org, make_plan):
        make_plan(org, name='Mensual')
        make_plan(org, name='Anual', duration_value=1, duration_unit='year')
        resp = admin_client.get(self.url)
        assert resp.status_code == status.HTTP_200_OK
        assert len(resp.data) == 2

    def test_manager_cannot_create_plan(self, api_client, make_user, org, make_membership):
        user = make_user(username='mgr@ex.com', email='mgr@ex.com', password='p!')
        make_membership(user, org, role='manager')
        token = _make_jwt_for(user, org=org, role='manager')
        api_client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')
        resp = api_client.post(self.url, {'name': 'X', 'price': '10'}, format='json')
        assert resp.status_code == status.HTTP_403_FORBIDDEN

    def test_org_without_payments_module_returns_403(self, api_client, make_user, make_org, make_membership):
        org = make_org(slug='no-billing', enabled_modules=['workers'])
        user = make_user(username='nb@ex.com', email='nb@ex.com', password='p!')
        make_membership(user, org, role='admin')
        token = _make_jwt_for(user, org=org, role='admin')
        api_client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')
        resp = api_client.get(self.url)
        assert resp.status_code == status.HTTP_403_FORBIDDEN

    def test_update_deactivates_plan(self, admin_client, org, make_plan):
        plan = make_plan(org)
        resp = admin_client.patch(f'{self.url}{plan.id}/', {'is_active': False}, format='json')
        assert resp.status_code == status.HTTP_200_OK
        assert resp.data['is_active'] is False

    def test_destroy_soft_deactivates(self, admin_client, org, make_plan):
        plan = make_plan(org)
        resp = admin_client.delete(f'{self.url}{plan.id}/')
        assert resp.status_code == status.HTTP_204_NO_CONTENT
        plan.refresh_from_db()
        assert plan.is_active is False


@pytest.mark.django_db
class TestRenewMembershipView:
    url = '/api/billing/renew/'

    def test_admin_can_renew(self, admin_client, org, make_member, make_plan):
        member = make_member(org)
        plan = make_plan(org)
        resp = admin_client.post(self.url, {
            'member_id': member.id, 'plan_id': plan.id, 'method': 'cash',
        }, format='json')
        assert resp.status_code == status.HTTP_201_CREATED
        assert resp.data['subscription']['status'] == 'active'
        assert resp.data['payment']['amount'] == str(plan.price)

    def test_invalid_plan_returns_404(self, admin_client, org, make_member):
        member = make_member(org)
        resp = admin_client.post(self.url, {
            'member_id': member.id, 'plan_id': 999999, 'method': 'cash',
        }, format='json')
        assert resp.status_code == status.HTTP_404_NOT_FOUND


@pytest.mark.django_db
class TestFreezeResumeViews:
    def test_freeze_then_resume(self, admin_client, org, make_member, make_plan):
        member = make_member(org)
        plan = make_plan(org)
        renew_resp = admin_client.post('/api/billing/renew/', {
            'member_id': member.id, 'plan_id': plan.id, 'method': 'cash',
        }, format='json')
        sub_id = renew_resp.data['subscription']['id']

        freeze_resp = admin_client.post(f'/api/billing/subscriptions/{sub_id}/freeze/')
        assert freeze_resp.status_code == status.HTTP_200_OK
        assert freeze_resp.data['status'] == 'frozen'

        resume_resp = admin_client.post(f'/api/billing/subscriptions/{sub_id}/resume/')
        assert resume_resp.status_code == status.HTTP_200_OK
        assert resume_resp.data['status'] == 'active'


@pytest.mark.django_db
class TestMemberBillingView:
    def test_returns_current_subscription_and_history(self, admin_client, org, make_member, make_plan):
        member = make_member(org)
        plan = make_plan(org)
        admin_client.post('/api/billing/renew/', {
            'member_id': member.id, 'plan_id': plan.id, 'method': 'cash',
        }, format='json')

        resp = admin_client.get(f'/api/billing/members/{member.id}/')
        assert resp.status_code == status.HTTP_200_OK
        assert resp.data['current_subscription']['status'] == 'active'
        assert len(resp.data['subscriptions']) == 1
        assert len(resp.data['payments']) == 1

    def test_member_without_subscription_returns_null(self, admin_client, org, make_member):
        member = make_member(org)
        resp = admin_client.get(f'/api/billing/members/{member.id}/')
        assert resp.status_code == status.HTTP_200_OK
        assert resp.data['current_subscription'] is None


@pytest.mark.django_db
class TestExpiringSubscriptionsView:
    url = '/api/billing/expiring/'

    def test_returns_subscriptions_within_window(self, admin_client, org, make_member, make_plan):
        member = make_member(org)
        plan = make_plan(org, duration_value=1, duration_unit='day')
        admin_client.post('/api/billing/renew/', {
            'member_id': member.id, 'plan_id': plan.id, 'method': 'cash',
        }, format='json')

        resp = admin_client.get(f'{self.url}?days=7')
        assert resp.status_code == status.HTTP_200_OK
        assert len(resp.data) == 1

    def test_far_future_subscription_not_included(self, admin_client, org, make_member, make_plan):
        member = make_member(org)
        plan = make_plan(org, duration_value=1, duration_unit='year')
        admin_client.post('/api/billing/renew/', {
            'member_id': member.id, 'plan_id': plan.id, 'method': 'cash',
        }, format='json')

        resp = admin_client.get(f'{self.url}?days=7')
        assert resp.status_code == status.HTTP_200_OK
        assert len(resp.data) == 0


@pytest.mark.django_db
class TestPaymentListView:
    url = '/api/billing/payments/'

    def test_lists_all_org_payments_with_registered_by_name(self, admin_client, admin_user, org, make_member, make_plan):
        member = make_member(org)
        plan = make_plan(org)
        resp = admin_client.post('/api/billing/renew/', {
            'member_id': member.id, 'plan_id': plan.id, 'method': 'cash',
        }, format='json')
        assert resp.status_code == status.HTTP_201_CREATED

        list_resp = admin_client.get(self.url)
        assert list_resp.status_code == status.HTTP_200_OK
        assert len(list_resp.data) == 1
        assert list_resp.data[0]['member_id'] == member.id
        assert list_resp.data[0]['member_name'] == f'{member.first_name} {member.last_name}'
        assert list_resp.data[0]['registered_by_id'] == admin_user.id
        assert list_resp.data[0]['registered_by_name']

    def test_filter_by_member_id(self, admin_client, org, make_member, make_plan):
        member_a = make_member(org, first_name='Ana', last_name='Uno')
        member_b = make_member(org, first_name='Bea', last_name='Dos')
        plan = make_plan(org)
        admin_client.post('/api/billing/renew/', {'member_id': member_a.id, 'plan_id': plan.id, 'method': 'cash'}, format='json')
        admin_client.post('/api/billing/renew/', {'member_id': member_b.id, 'plan_id': plan.id, 'method': 'card'}, format='json')

        resp = admin_client.get(f'{self.url}?member_id={member_a.id}')
        assert resp.status_code == status.HTTP_200_OK
        assert len(resp.data) == 1
        assert resp.data[0]['member_id'] == member_a.id

    def test_filter_by_method(self, admin_client, org, make_member, make_plan):
        member = make_member(org)
        plan = make_plan(org)
        admin_client.post('/api/billing/renew/', {'member_id': member.id, 'plan_id': plan.id, 'method': 'transfer'}, format='json')

        resp = admin_client.get(f'{self.url}?method=transfer')
        assert resp.status_code == status.HTTP_200_OK
        assert len(resp.data) == 1

        resp_other = admin_client.get(f'{self.url}?method=cash')
        assert resp_other.status_code == status.HTTP_200_OK
        assert len(resp_other.data) == 0

    def test_filter_by_date_range(self, admin_client, org, make_member, make_plan):
        from datetime import timedelta

        from django.utils import timezone

        from apps.billing.models import DuesPaymentModel

        member = make_member(org)
        DuesPaymentModel.objects.create(
            organization=org, member=member, amount='50.00', method='cash',
            paid_at=timezone.localdate() - timedelta(days=10),
        )
        DuesPaymentModel.objects.create(
            organization=org, member=member, amount='60.00', method='cash',
            paid_at=timezone.localdate(),
        )

        today = timezone.localdate().isoformat()
        resp = admin_client.get(f'{self.url}?date_from={today}&date_to={today}')
        assert resp.status_code == status.HTTP_200_OK
        assert len(resp.data) == 1
        assert resp.data[0]['amount'] == '60.00'

    def test_org_scoped_only(self, admin_client, org, make_org, make_member, make_plan):
        other_org = make_org(slug='other-payments-org')
        other_member = make_member(other_org)
        plan = make_plan(other_org)
        # Renovar en otra organización no debe aparecer en el historial de `org`
        from apps.billing.models import DuesPaymentModel
        DuesPaymentModel.objects.create(
            organization=other_org, member=other_member, amount='99.00', method='cash', paid_at='2026-01-01',
        )

        resp = admin_client.get(self.url)
        assert resp.status_code == status.HTTP_200_OK
        assert len(resp.data) == 0

    def test_org_without_payments_module_returns_403(self, api_client, make_user, make_org, make_membership):
        org = make_org(slug='no-billing-payments', enabled_modules=['workers'])
        user = make_user(username='nbp@ex.com', email='nbp@ex.com', password='p!')
        make_membership(user, org, role='admin')
        token = _make_jwt_for(user, org=org, role='admin')
        api_client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')
        resp = api_client.get(self.url)
        assert resp.status_code == status.HTTP_403_FORBIDDEN


@pytest.mark.django_db
class TestDailyStatsReportView:
    url = '/api/billing/reports/daily/'

    def test_default_range_includes_today(self, admin_client):
        resp = admin_client.get(self.url)
        assert resp.status_code == status.HTTP_200_OK
        assert len(resp.data) == 30
        assert resp.data[-1]['date'] == timezone.localdate().isoformat()

    def test_includes_a_backdated_payment(self, admin_client, org, make_member):
        from apps.billing.models import DuesPaymentModel

        yesterday = timezone.localdate() - timedelta(days=1)
        member = make_member(org)
        DuesPaymentModel.objects.create(
            organization=org, member=member, amount='75.00', method='cash', paid_at=yesterday,
        )

        resp = admin_client.get(f'{self.url}?date_from={yesterday.isoformat()}&date_to={yesterday.isoformat()}')
        assert resp.status_code == status.HTTP_200_OK
        assert len(resp.data) == 1
        assert resp.data[0]['revenue'] == '75.00'

    def test_org_without_payments_module_returns_403(self, api_client, make_user, make_org, make_membership):
        org = make_org(slug='no-billing-daily', enabled_modules=['workers'])
        user = make_user(username='nbd@ex.com', email='nbd@ex.com', password='p!')
        make_membership(user, org, role='admin')
        token = _make_jwt_for(user, org=org, role='admin')
        api_client.credentials(HTTP_AUTHORIZATION=f'Bearer {token}')
        resp = api_client.get(self.url)
        assert resp.status_code == status.HTTP_403_FORBIDDEN


@pytest.mark.django_db
class TestDailyStatsExportView:
    url = '/api/billing/reports/export/'

    def test_xlsx_export_returns_valid_workbook(self, admin_client, org, make_member, make_plan):
        from io import BytesIO
        from openpyxl import load_workbook

        member = make_member(org)
        plan = make_plan(org)
        admin_client.post('/api/billing/renew/', {
            'member_id': member.id, 'plan_id': plan.id, 'method': 'cash',
        }, format='json')

        resp = admin_client.get(f'{self.url}?export_format=xlsx')
        assert resp.status_code == status.HTTP_200_OK
        assert resp['Content-Type'] == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        assert 'attachment' in resp['Content-Disposition']

        wb = load_workbook(BytesIO(resp.content))
        assert wb.sheetnames == ['Diario', 'Ingresos por mes', 'Estado de membresías', 'Asistencia por día']

    def test_pdf_export_returns_pdf_bytes(self, admin_client):
        resp = admin_client.get(f'{self.url}?export_format=pdf')
        assert resp.status_code == status.HTTP_200_OK
        assert resp['Content-Type'] == 'application/pdf'
        assert resp.content.startswith(b'%PDF')

    def test_invalid_format_returns_400(self, admin_client):
        resp = admin_client.get(f'{self.url}?export_format=docx')
        assert resp.status_code == status.HTTP_400_BAD_REQUEST
